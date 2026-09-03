"""
Uji aturan sel harian laporan (1 / 0 / '-' / kosong) — dikonfirmasi pengguna
2026-09-03 — pada tiga keluaran: Excel rekap, Excel bulanan, preview HTML.

Jalankan:  python manage.py test tests.test_report_cells
"""
import datetime
import os
import shutil
import tempfile
from unittest import mock

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from core.ingest import ingest_uploaded_files
from core.models import Project
from core.parsers.ghp import detect_ghp_range, parse_ghp
from core.report import (
    CELL_OUTSIDE, _find_flight_headers, _format_day_of_flight, _month_row_period,
    _parse_day_of_flight, _parse_periode_range, generate_report, generate_rekap_report,
    get_project_report_data, resolve_daily_cell,
)

SRC = os.path.join('docs', 'Source')
WTT_MARET = os.path.join(SRC, 'Working Time Table [WTT] Maret 2026.pdf')
WTT_APRIL = os.path.join(SRC, 'Working Time Table [WTT] April 2026.pdf')
WTT_JULI = os.path.join(SRC, 'Working Time Table [WTT] Juli 2026 (1).pdf')
GHP_MARET = os.path.join(SRC, 'data_master_maret_2026.xls')
GHP_APRIL = os.path.join(SRC, 'data_master_april_2026.xls')
GHP_JULI = os.path.join(SRC, 'data_master_juli_2026.xls')
PPRP_BTH = os.path.join(SRC, 'PPRP BTH-SUB S26 UPDATE.pdf')   # QG948, berlaku 13 Jul 2026
TEMPLATE = os.path.join('static', 'tpl', 'form_realisasi_winter26.xlsx')

D = datetime.date


class DailyCellRuleTests(SimpleTestCase):
    """Fungsi murni — tanpa database."""

    def test_day_of_flight_notations(self):
        self.assertEqual(_parse_day_of_flight('1234567'), frozenset({1, 2, 3, 4, 5, 6, 7}))
        self.assertEqual(_parse_day_of_flight('1004507'), frozenset({1, 4, 5, 7}))
        self.assertEqual(_parse_day_of_flight(204060), frozenset({2, 4, 6}), 'nol depan hilang di Excel')
        self.assertEqual(_parse_day_of_flight(1234567.0), frozenset({1, 2, 3, 4, 5, 6, 7}))
        self.assertEqual(_parse_day_of_flight('246'), frozenset({2, 4, 6}), 'notasi ringkas surat PPRP')
        self.assertEqual(_parse_day_of_flight('1--45-7'), frozenset({1, 4, 5, 7}))
        self.assertIsNone(_parse_day_of_flight(None))
        self.assertIsNone(_parse_day_of_flight('  '))
        self.assertIsNone(_parse_day_of_flight('Perpanjangan'))

    def test_day_of_flight_round_trip(self):
        self.assertEqual(_format_day_of_flight(frozenset({2, 4, 6})), '0204060')
        self.assertEqual(_format_day_of_flight(None), '1234567')
        for s in ('1234567', '1004507', '0204060'):
            self.assertEqual(_format_day_of_flight(_parse_day_of_flight(s)), s)

    def test_periode_parsing_tolerates_template_variants(self):
        self.assertEqual(_parse_periode_range('29 MAR 2026/24 OKT 2026'), (D(2026, 3, 29), D(2026, 10, 24)))
        self.assertEqual(_parse_periode_range('11 APR 2026/24OKT2026'), (D(2026, 4, 11), D(2026, 10, 24)))
        self.assertEqual(_parse_periode_range('1 MAR 2026/12 JUL 2026'), (D(2026, 3, 1), D(2026, 7, 12)))
        self.assertIsNone(_parse_periode_range('PERLU REVIEW - SURAT TIDAK TERBACA'))
        self.assertIsNone(_parse_periode_range(None))

    def test_month_row_period(self):
        self.assertEqual(_month_row_period(datetime.datetime(2026, 4, 4), 2026), (2026, 4))
        self.assertEqual(_month_row_period('Mei-26', 2026), (2026, 5))
        self.assertEqual(_month_row_period('Sep-26', 2027), (2026, 9))
        self.assertEqual(_month_row_period('Oktober', 2026), (2026, 10))
        self.assertIsNone(_month_row_period(None, 2026))

    def test_resolve_order_of_rules(self):
        ghp = {'flown': {('QG179', D(2026, 7, 14))}, 'coverage': [(D(2026, 7, 1), D(2026, 7, 24))]}
        periode = (D(2026, 4, 11), D(2026, 10, 24))
        days = frozenset({2, 4, 6})
        # 1) di luar periode
        self.assertEqual(resolve_daily_cell(D(2026, 4, 10), periode, days, ghp, 'QG179'), CELL_OUTSIDE)
        # 2) bukan hari operasi (13 Jul 2026 = Senin)
        self.assertEqual(resolve_daily_cell(D(2026, 7, 13), periode, days, ghp, 'QG179'), CELL_OUTSIDE)
        # 3) di luar cakupan GHP (Selasa 28 Jul, file hanya sampai 24)
        self.assertIsNone(resolve_daily_cell(D(2026, 7, 28), periode, days, ghp, 'QG179'))
        # 4) terbang
        self.assertEqual(resolve_daily_cell(D(2026, 7, 14), periode, days, ghp, 'QG179'), 1)
        # 5) direncanakan, tidak terbang (Kamis 16 Jul)
        self.assertEqual(resolve_daily_cell(D(2026, 7, 16), periode, days, ghp, 'QG179'), 0)
        # periode/pola tak diketahui -> tidak pernah '-' karena tebakan
        self.assertEqual(resolve_daily_cell(D(2026, 7, 13), None, None, ghp, 'QG179'), 0)

    def test_ghp_range_from_header(self):
        self.assertEqual(detect_ghp_range(GHP_MARET), (D(2026, 3, 1), D(2026, 3, 31)))

    def test_block_month_rows_repairs_a_mislabelled_row_by_position(self):
        """Template resmi: blok QG-672 baris pertamanya berlabel Okt-26 (salah ketik) — dibaca sebagai Maret."""
        from openpyxl import Workbook
        from core.report import _block_month_rows
        ws = Workbook().active
        labels = [datetime.datetime(2026, 10, 1), datetime.datetime(2026, 4, 4), 'Mei-26'] + \
                 [datetime.datetime(2026, m, 1) for m in (6, 7, 8, 9, 10)]
        for i, v in enumerate(labels, start=5):
            ws.cell(i, 13).value = v
        self.assertEqual([m for _, (_, m) in _block_month_rows(ws, 5, 12, 2026)], [3, 4, 5, 6, 7, 8, 9, 10])
        # Blok yang labelnya benar atau tidak lengkap (QG-179: Apr..Sep) dibiarkan apa adanya
        ws2 = Workbook().active
        for i, m in enumerate((4, 5, 6, 7, 8, 9), start=5):
            ws2.cell(i, 13).value = datetime.datetime(2026, m, 1)
        self.assertEqual([m for _, (_, m) in _block_month_rows(ws2, 5, 10, 2026)], [4, 5, 6, 7, 8, 9])


def upload(path):
    with open(path, 'rb') as fh:
        return SimpleUploadedFile(os.path.basename(path), fh.read())


class ReportCellsTests(TestCase):
    """End-to-end: Maret + April + Juli 2026 dari docs/Source, lalu periksa sel tertentu."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('tester', 'tester@example.com', 'x', is_staff=True)
        cls.upload_dir = tempfile.mkdtemp(prefix='report-test-')
        cls.out_dir = tempfile.mkdtemp(prefix='report-out-')

        def ingest(file_type, path, project=None):
            res = ingest_uploaded_files(file_type, [upload(path)], cls.user, project=project,
                                        upload_dir=cls.upload_dir)
            assert res.success, [f.message for f in res.files]
            return res.project

        cls.maret = ingest('WTT', WTT_MARET)
        ingest('GHP', GHP_MARET, cls.maret)
        cls.april = ingest('WTT', WTT_APRIL)
        ingest('GHP', GHP_APRIL, cls.april)
        cls.juli = ingest('WTT', WTT_JULI)
        ingest('PPRP', PPRP_BTH, cls.juli)
        ingest('GHP', GHP_JULI, cls.juli)
        for p in (cls.maret, cls.april, cls.juli):
            p.template_path = TEMPLATE
            p.save()

        cls.rekap_path = os.path.join(cls.out_dir, 'rekap.xlsx')
        generate_rekap_report(cls.rekap_path)
        cls.rekap = _Sheet(cls.rekap_path)

        # Kebenaran GHP langsung dari file (pembanding independen dari kode laporan)
        cls.flew = set()
        for path, year in ((GHP_MARET, 2026), (GHP_APRIL, 2026), (GHP_JULI, 2026)):
            for r in parse_ghp(path, year=year):
                if r['origin'] == 'SUB':
                    cls.flew.add((r['flight_number'], datetime.date.fromisoformat(r['flight_date'])))

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls.upload_dir, ignore_errors=True)
        shutil.rmtree(cls.out_dir, ignore_errors=True)

    def expect_ghp(self, flight, day):
        return 1 if (flight, day) in self.flew else 0

    # --- aturan 1: di luar periode blok ---

    def test_before_season_start_is_outside(self):
        """Template: QG-430 berizin 29 MAR-24 OKT -> 1-28 Maret di luar jadwal."""
        for day in (1, 15, 28):
            self.assertEqual(self.rekap.cell('QG430', 0, 3, day), CELL_OUTSIDE)
        self.assertEqual(self.rekap.cell('QG430', 0, 3, 29), self.expect_ghp('QG430', D(2026, 3, 29)))

    def test_segments_do_not_bleed_into_each_other(self):
        """QG948: blok SEMULA berakhir 12 Jul, blok surat mulai 13 Jul."""
        self.assertEqual(self.rekap.block_count('QG948'), 2)
        self.assertIn('12 JUL 2026', self.rekap.periode('QG948', 0))
        self.assertTrue(self.rekap.periode('QG948', 1).startswith('13 JUL 2026'))
        self.assertEqual(self.rekap.cell('QG948', 0, 7, 12), self.expect_ghp('QG948', D(2026, 7, 12)))
        self.assertEqual(self.rekap.cell('QG948', 0, 7, 13), CELL_OUTSIDE)
        self.assertEqual(self.rekap.cell('QG948', 1, 7, 12), CELL_OUTSIDE)
        self.assertEqual(self.rekap.cell('QG948', 1, 7, 13), self.expect_ghp('QG948', D(2026, 7, 13)))
        self.assertEqual(self.rekap.cell('QG948', 1, 3, 30), CELL_OUTSIDE, 'Maret di luar periode surat')

    # --- aturan 2: pola hari ---

    def test_non_operating_weekdays_are_outside_not_zero(self):
        """QG-179 (Day Of Flight 0204060 = Sel/Kam/Sab): Senin '-', Selasa 1/0."""
        self.assertEqual(self.rekap.day_of_flight('QG179', 0), frozenset({2, 4, 6}))
        self.assertEqual(self.rekap.cell('QG179', 0, 7, 13), CELL_OUTSIDE)   # Senin
        self.assertEqual(self.rekap.cell('QG179', 0, 7, 15), CELL_OUTSIDE)   # Rabu
        self.assertEqual(self.rekap.cell('QG179', 0, 7, 14), self.expect_ghp('QG179', D(2026, 7, 14)))  # Selasa
        self.assertEqual(self.rekap.cell('QG179', 0, 4, 10), CELL_OUTSIDE, 'sebelum 11 APR (periode template)')
        self.assertEqual(self.rekap.cell('QG179', 0, 4, 14), self.expect_ghp('QG179', D(2026, 4, 14)))

    # --- aturan 3: cakupan GHP ---

    def test_months_without_ghp_are_blank_not_zero(self):
        for month in (5, 6, 8, 9, 10):
            self.assertIsNone(self.rekap.cell('QG430', 0, month, 5), f'bulan {month} tanpa GHP harus kosong')

    # --- aturan 4/5: realisasi dari GHP, termasuk flight tanpa baris WTT ---

    def test_permitted_flight_missing_from_wtt_still_gets_actuals(self):
        """QG719 tidak ada di PDF WTT Maret, tapi berizin di template dan terbang menurut GHP."""
        self.assertEqual(self.rekap.cell('QG719', 0, 3, 28), CELL_OUTSIDE)
        for day in (29, 30, 31):
            self.assertEqual(self.rekap.cell('QG719', 0, 3, day), self.expect_ghp('QG719', D(2026, 3, day)))
        self.assertIn(1, [self.rekap.cell('QG719', 0, 3, d) for d in (29, 30, 31)])

    def test_early_april_outside_wtt_but_inside_permit_is_actual(self):
        """WTT April mulai tgl 7; 1-6 April tetap dalam izin (29 MAR-24 OKT) dan GHP mencatat terbang."""
        for day in range(1, 7):
            self.assertEqual(self.rekap.cell('QG430', 0, 4, day), self.expect_ghp('QG430', D(2026, 4, day)))
        self.assertEqual(sum(self.rekap.cell('QG430', 0, 4, d) for d in range(1, 7)), 6)

    def test_every_cell_matches_ghp_truth_for_july(self):
        """Semua blok, seluruh Juli: 1/0 harus persis isi file GHP, tidak ada 0 tanpa rencana."""
        for flight in ('QG430', 'QG725', 'QG692', 'QG723'):
            for day in range(1, 32):
                v = self.rekap.cell(flight, 0, 7, day)
                if v in (0, 1):
                    self.assertEqual(v, self.expect_ghp(flight, D(2026, 7, day)), f'{flight} {day} Jul')

    def test_totals_formula_survives(self):
        self.assertTrue(str(self.rekap.total_formula('QG430', 0)).startswith('=SUM('))

    # --- laporan bulanan & preview memakai aturan yang sama ---

    def test_monthly_report_leaves_other_months_blank(self):
        path = os.path.join(self.out_dir, 'bulanan_juli.xlsx')
        generate_report(self.juli.id, TEMPLATE, path)
        sheet = _Sheet(path)
        self.assertIsNone(sheet.cell('QG430', 0, 3, 30), 'Maret tidak dicakup GHP project Juli -> kosong')
        self.assertIsNone(sheet.cell('QG430', 0, 4, 10))
        self.assertEqual(sheet.cell('QG430', 0, 7, 10), self.expect_ghp('QG430', D(2026, 7, 10)))
        self.assertEqual(sheet.cell('QG430', 0, 3, 10), CELL_OUTSIDE, "di luar periode tetap '-' meski tanpa GHP")

    def test_preview_matches_excel_cells(self):
        data = get_project_report_data(self.juli.id)
        rows = {}
        for r in data['rows']:
            rows.setdefault(r['flight_number'], r)   # baris pertama per flight = SEMULA
        qg179 = rows['QG-179']
        self.assertEqual(qg179['day_pattern'], '0204060')
        self.assertEqual(qg179['days'][12]['val'], CELL_OUTSIDE)     # 13 Jul Senin
        self.assertEqual(qg179['days'][13]['val'], str(self.expect_ghp('QG179', D(2026, 7, 14))))
        for flight in ('QG-430', 'QG-723'):
            for i, cell in enumerate(rows[flight]['days']):
                excel = self.rekap.cell(flight.replace('-', ''), 0, 7, i + 1)
                self.assertEqual(cell['val'], '' if excel is None else str(excel), f'{flight} hari {i + 1}')

    def test_stale_charter_rows_in_db_never_reach_the_report(self):
        """Baris charter yang dimuat kode lama masih bisa ada di DB; laporan harus mengabaikannya."""
        from core.models import ScheduleVersion
        ScheduleVersion.objects.create(
            project=self.maret, flight_number='QG9177', origin='SUB', destination='HLP',
            flight_date=D(2026, 3, 19), std=datetime.time(8, 35), sta=datetime.time(10, 5), is_active=True)

        data = get_project_report_data(self.maret.id)

        self.assertNotIn('QG-9177', [r['flight_number'] for r in data['rows']])
        self.assertNotIn('QG9177', [r['flight_number'] for r in data['rows']])

    # --- ketiga keluaran harus memuat data yang sama (kepentingan audit Otban) ---

    def test_monthly_excel_equals_preview_in_every_column(self):
        path = os.path.join(self.out_dir, 'bulanan_juli_cmp.xlsx')
        generate_report(self.juli.id, TEMPLATE, path)
        excel = _blocks_from_excel(path, 7)
        preview = _blocks_from_preview(self.juli.id)

        self.assertEqual(set(excel), set(preview), 'daftar blok (flight, urutan) harus sama')
        self.assertEqual(excel[('QG948', 0)]['tipe'], 'Perpanjangan')
        self.assertEqual(excel[('QG948', 1)]['surat'], preview[('QG948', 1)]['surat'])
        for key in excel:
            for field in ('etd', 'eta', 'atd', 'ata', 'periode', 'surat', 'tipe', 'hari', 'days'):
                self.assertEqual(excel[key][field], preview[key][field], f'{key} kolom {field}')

    def test_monthly_excel_is_a_slice_of_the_rekap(self):
        path = os.path.join(self.out_dir, 'bulanan_juli_slice.xlsx')
        generate_report(self.juli.id, TEMPLATE, path)
        monthly = _blocks_from_excel(path, 7)
        rekap = _blocks_from_excel(self.rekap_path, 7)

        self.assertEqual(set(monthly), set(rekap))
        for key in monthly:
            for field in ('etd', 'eta', 'atd', 'ata', 'periode', 'surat', 'tipe', 'hari', 'days'):
                self.assertEqual(monthly[key][field], rekap[key][field], f'{key} kolom {field}')

    def test_template_submission_type_is_preserved(self):
        """Template menulis 'Perubahan' untuk QG-179; laporan tidak boleh menimpanya."""
        self.assertEqual(self.rekap.tipe('QG179', 0), 'Perubahan')
        self.assertEqual(self.rekap.tipe('QG430', 0), 'Perpanjangan')

    def test_day_of_flight_written_with_seven_characters(self):
        self.assertEqual(self.rekap.raw_day_of_flight('QG179', 0), '0204060')
        self.assertEqual(self.rekap.raw_day_of_flight('QG723', 0), '1004507')

    def test_blank_template_row_falls_back_to_wtt_without_inventing_a_letter(self):
        """Musim baru: template bisa hanya berisi nomor flight & rute. Periode dan
        Day Of Flight diambil dari WTT dan ditulis ke Excel; surat dibiarkan kosong."""
        blank_tpl = os.path.join(self.out_dir, 'template_kosong.xlsx')
        wb = load_workbook(TEMPLATE)
        ws = wb.active
        for r in range(8, ws.max_row + 1):
            if str(ws.cell(r, 2).value).strip() == 'QG-430':
                for col in (4, 5, 6, 7, 8, 9, 10, 11):
                    ws.cell(r, col).value = None
                break
        wb.save(blank_tpl)

        path = os.path.join(self.out_dir, 'bulanan_juli_blank.xlsx')
        generate_report(self.juli.id, blank_tpl, path)
        sheet = _Sheet(path)
        self.assertEqual(sheet.periode('QG430', 0), '1 MAR 2026/31 JUL 2026')   # rentang WTT yang dimuat
        self.assertEqual(sheet.raw_day_of_flight('QG430', 0), '1234567')
        self.assertIn(str(self.ws_cell(path, 'QG430', 9)), ('None', ''))         # surat tidak dikarang
        for day in (1, 10, 31):
            self.assertEqual(sheet.cell('QG430', 0, 7, day), self.expect_ghp('QG430', D(2026, 7, day)))
        # 1 Maret ada di rentang WTT (bukan '-'), tapi GHP project Juli tidak
        # mencakupnya -> kosong, bukan 0 dan bukan '-'.
        self.assertIsNone(sheet.cell('QG430', 0, 3, 1))

        self.juli.template_path = blank_tpl
        self.juli.save()
        row = next(r for r in get_project_report_data(self.juli.id)['rows'] if r['flight_number'] == 'QG-430')
        self.assertEqual(row['periode'], '1 MAR 2026/31 JUL 2026')
        self.assertEqual(row['day_pattern'], '1234567')
        self.assertEqual(row['pprp_no'], '-')

    @staticmethod
    def ws_cell(path, flight, col):
        ws = load_workbook(path).active
        for r in range(8, ws.max_row + 1):
            if str(ws.cell(r, 2).value).replace('-', '').strip() == flight:
                return ws.cell(r, col).value
        return None

    def test_preview_lists_permitted_flight_absent_from_wtt(self):
        data = get_project_report_data(self.maret.id)
        qg719 = [r for r in data['rows'] if r['flight_number'] == 'QG-719']
        self.assertEqual(len(qg719), 1)
        # Kebalikannya: ada di WTT tapi tidak di template & tanpa surat (QG715,
        # musim sebelumnya) -> tidak masuk laporan musim ini, seperti di Excel.
        self.assertNotIn('QG-715', [r['flight_number'] for r in data['rows']])
        self.assertEqual(qg719[0]['days'][27]['val'], CELL_OUTSIDE)   # 28 Mar
        self.assertEqual(qg719[0]['days'][29]['val'], str(self.expect_ghp('QG719', D(2026, 3, 30))))


class _Sheet:
    """Akses sel laporan per (flight, indeks blok, bulan, tanggal)."""

    def __init__(self, path):
        self.ws = load_workbook(path).active
        from core.report import _detect_columns
        self.cols, _ = _detect_columns(self.ws)
        self.blocks = {}
        for row_start, norm, _ in _find_flight_headers(self.ws, self.cols['flight']):
            self.blocks.setdefault(norm, []).append(row_start)

    def block_count(self, flight):
        return len(self.blocks.get(flight, []))

    def _top(self, flight, block):
        return self.blocks[flight][block]

    def _month_row(self, flight, block, month):
        top = self._top(flight, block)
        for r in range(top, top + 12):
            v = _month_row_period(self.ws.cell(r, self.cols.get('bulan_tahun', 13)).value, 2026)
            if v and v[1] == month:
                return r
        raise AssertionError(f'baris bulan {month} tidak ditemukan untuk {flight} blok {block}')

    def cell(self, flight, block, month, day):
        return self.ws.cell(self._month_row(flight, block, month), self.cols['day_start'] + day - 1).value

    def periode(self, flight, block):
        return str(self.ws.cell(self._top(flight, block), self.cols['periode']).value)

    def day_of_flight(self, flight, block):
        return _parse_day_of_flight(self.ws.cell(self._top(flight, block), self.cols['hari']).value)

    def total_formula(self, flight, block):
        return self.ws.cell(self._month_row(flight, block, 7), self.cols['total']).value

    def tipe(self, flight, block):
        return str(self.ws.cell(self._top(flight, block), self.cols['tipe']).value).strip()

    def raw_day_of_flight(self, flight, block):
        return str(self.ws.cell(self._top(flight, block), self.cols['hari']).value)


def _norm_time(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.time, datetime.datetime)):
        return v.strftime('%H:%M')
    return str(v).strip()


def _blocks_from_excel(path, month):
    """{(flight, urutan blok): kolom-kolom + 31 sel bulan `month`} dari file Excel.
    Blok tanpa baris bulan itu (mis. QG-179 tidak punya baris Maret) -> 31 x '-'."""
    from collections import Counter
    from core.report import _detect_columns, _normalize_flight
    ws = load_workbook(path).active
    cols, _ = _detect_columns(ws)
    headers = _find_flight_headers(ws, cols['flight'])
    out, seen = {}, Counter()
    for i, (top, norm, _) in enumerate(headers):
        nxt = headers[i + 1][0] if i + 1 < len(headers) else ws.max_row + 1
        mrow = next((r for r in range(top, nxt)
                     if (_month_row_period(ws.cell(r, 13).value, 2026) or (0, 0))[1] == month), None)
        if mrow:
            days = ['' if v is None else str(v) for v in
                    (ws.cell(mrow, cols['day_start'] + d).value for d in range(31))]
        else:
            days = [CELL_OUTSIDE] * 31
        out[(norm, seen[norm])] = {
            'etd': _norm_time(ws.cell(top, cols['etd']).value), 'eta': _norm_time(ws.cell(top, cols['eta']).value),
            'atd': _norm_time(ws.cell(top, cols['atd']).value), 'ata': _norm_time(ws.cell(top, cols['ata']).value),
            'periode': _norm_time(ws.cell(top, cols['periode']).value),
            'surat': _norm_time(ws.cell(top, cols['surat']).value) or '-',
            'tipe': _norm_time(ws.cell(top, cols['tipe']).value),
            'hari': _norm_time(ws.cell(top, cols['hari']).value),
            'days': days,
        }
        seen[norm] += 1
    return out


def _blocks_from_preview(project_id):
    from collections import Counter
    from core.report import _normalize_flight
    out, seen = {}, Counter()
    for r in get_project_report_data(project_id)['rows']:
        norm = _normalize_flight(r['flight_number'])
        out[(norm, seen[norm])] = {
            'etd': _norm_time(r['etd']), 'eta': _norm_time(r['eta']),
            'atd': _norm_time(r['atd']), 'ata': _norm_time(r['ata']),
            'periode': _norm_time(r['periode']), 'surat': _norm_time(r['pprp_no']) or '-',
            'tipe': _norm_time(r['pprp_type']), 'hari': _norm_time(r['day_pattern']),
            'days': [c['val'] for c in r['days']],
        }
        seen[norm] += 1
    return out
